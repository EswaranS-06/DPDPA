#!/usr/bin/env python3
"""DUATF Master Register (Excel) - same IDs as the Obsidian vault."""
import sys, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

from data_legal import *
from data_obligations import DOMAINS, OBLIGATIONS
from data_controls import C as CONTROLS
from data_processes_common import COMMON
from data_sectors import SECTORS
from data_vocab import *
import data_playbook as PB

OUT = "/home/claude/out/DPDP-DUATF-Master-Register.xlsx"
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
INPUT = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE = Font(name=F, italic=True, color="0000FF", size=10)
BODY = Font(name=F, size=10)
TITLE = Font(name=F, bold=True, size=14, color="1F3864")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, headers, rows, widths=None, input_cols=(), start=4, note=None, blank_rows=0, example_rows=0):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=start + 1 + i, column=j, value=v)
            c.font = EXAMPLE if i < example_rows else BODY
            c.alignment, c.border = WRAP, BORDER
            if j - 1 in input_cols:
                c.fill = INPUT
    for i in range(len(rows), len(rows) + blank_rows):
        for j in range(1, len(headers) + 1):
            c = ws.cell(row=start + 1 + i, column=j)
            c.border, c.font, c.alignment = BORDER, BODY, WRAP
            if j - 1 in input_cols:
                c.fill = INPUT
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[CL(j)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f"A{start}:{CL(len(headers))}{start + max(len(rows) + blank_rows, 1)}"
    ws.sheet_view.zoomScale = 90
    return ws

def dv_list(ws, rng, src):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)

pen_txt = {k: v[0] + (f" (Rs {v[1]} cr)" if v[1] and v[1] >= 1 else "") for k, v in PENALTIES.items()}
dom_name = {d[0]: d[1] for d in DOMAINS}
ctl_for = {}
for c in CONTROLS:
    for o in c["obligations"]:
        ctl_for.setdefault(o, []).append(c["id"])

# ---------------------------------------------------------------- README
ws = wb.create_sheet("README")
lines = [
 ("DPDP Unified Assessment & Tracking Framework (DUATF) - Master Register", TITLE),
 ("Version 1.0 (24 Sep 2026). Companion to the Obsidian vault: same IDs (OBL-, CTL-, ACT-, FND- ...).", BODY),
 ("", BODY),
 ("HOW TO USE", Font(name=F, bold=True, size=11)),
 ("1. Reference tabs (blue headers, no yellow): Obligations, Controls, Obl-Ctl Map, Act, Rules & Schedules, Sectors, Retention Anchors, Process Catalogue, Data Elements, Lists. Don't edit per client.", BODY),
 ("2. Working tabs (yellow cells = you fill in): Entity Profile, Activity Register, Control Tests, Findings, Remediation, Rights Log, Breach Register, Third Parties, Transfers, Evidence (PBC).", BODY),
 ("3. The first data row on each working tab is an EXAMPLE (blue italic) for the fictional DemoPay. Overwrite or delete it.", BODY),
 ("4. Activity Register: set lawful basis and flags as Y/N. The Applicability tab then computes, per obligation x activity, Y / N / N/A-17(1) for the first 30 activities.", BODY),
 ("5. Findings: Risk score = Likelihood x Impact; the band is automatic. Rights Log: due date = received + internal SLA (max 90 days). Breach Register: CERT-In 6h and Board 72h deadlines are computed from the awareness timestamp.", BODY),
 ("6. Dashboard summarises everything with formulas.", BODY),
 ("", BODY),
 ("KEY DATES", Font(name=F, bold=True, size=11)),
 ("Phase 1: 13 Nov 2025 (in force) - Board, definitions, Rules 1, 2, 17-21", BODY),
 ("Phase 2: 13 Nov 2026 (computed) - Consent Manager registration: s.6(9), s.27(1)(d), Rule 4", BODY),
 ("Phase 3: 13 May 2027 (computed) - all substantive obligations: Rules 3, 5-16, 22, 23", BODY),
 (DATE_NOTE, Font(name=F, italic=True, size=9)),
 ("", BODY),
 ("SOURCES", Font(name=F, bold=True, size=11)),
 ("DPDP Act 2023 (Act 22 of 2023); DPDP Rules 2025 G.S.R. 846(E); Commencement G.S.R. 843(E) - official Gazette text via egazette.gov.in / meity.gov.in; section and rule text cross-checked on dpdprules.org (verified Sept 2026). CERT-In Directions 28.04.2022. Sector instruments named per row; rows marked 'verify' must be checked against the current text before client use.", BODY),
 ("Not legal advice.", Font(name=F, italic=True, size=9)),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = f
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 150
ws.cell(row=len(lines) + 2, column=1, value="Legend:").font = Font(name=F, bold=True)
c = ws.cell(row=len(lines) + 3, column=1, value="Yellow cell = input")
c.fill, c.font = INPUT, BODY
c = ws.cell(row=len(lines) + 4, column=1, value="Blue italic = example row (replace)")
c.font = EXAMPLE

# ---------------------------------------------------------------- Lists (for validation)
lists = {
 "Basis": [b[0] for b in LAWFUL_BASES],
 "YN": ["Y", "N"],
 "Rating": ["0", "1", "2", "3", "4", "NA"],
 "Scale": ["1", "2", "3", "4", "5"],
 "FindingStatus": FINDING_STATUS,
 "TestType": [t[0] for t in TEST_TYPES],
 "Role": ENTITY_ROLES,
 "DigitalState": [d[0] for d in DIGITAL_STATES],
 "Hosting": HOSTING,
 "TPType": THIRD_PARTY_TYPES,
 "Right": ["access", "correction", "completion", "updating", "erasure", "nomination", "grievance", "consent withdrawal"],
 "Confidence": ["low", "medium", "high"],
 "Penalty": list(PENALTIES.keys()),
 "Transfer": TRANSFER_TYPES,
 "Channel": CHANNELS,
 "Domain": [f"{d[0]} {d[1]}" for d in DOMAINS],
 "EvidenceStatus": ["requested", "received", "insufficient", "accepted", "n/a"],
 "DPAStatus": ["none", "signed", "vendor_paper_accepted", "in negotiation", "n/a (independent DF)"],
}
wsL = wb.create_sheet("Lists")
wsL["A1"] = "Controlled vocabularies (used by drop-downs)"
wsL["A1"].font = TITLE
LREF = {}
for j, (k, vals) in enumerate(lists.items(), 1):
    wsL.cell(row=3, column=j, value=k).font = HDR_FONT
    wsL.cell(row=3, column=j).fill = HDR_FILL
    for i, v in enumerate(vals):
        wsL.cell(row=4 + i, column=j, value=v).font = BODY
    wsL.column_dimensions[CL(j)].width = 24
    LREF[k] = f"Lists!${CL(j)}$4:${CL(j)}${3 + len(vals)}"

# ---------------------------------------------------------------- Obligations reference
BASIS_COLS = ["consent", "s7a", "s7b", "s7c", "s7d", "s7e", "s7f", "s7g", "s7h", "s7i", "ex17_1", "ex17_2"]
FLAG_COLS = [f[0] for f in FLAGS]
def trig_basis(t):
    b = t.get("basis", [])
    out = []
    for x in b:
        if x.startswith("ex17_1"):
            x = "ex17_1"
        elif x.startswith("ex17_2"):
            x = "ex17_2"
        elif x == "ex17_3":
            continue
        if x not in out:
            out.append(x)
    return ",".join(out)
obl_rows = []
for (oid, dom, title, req, act, rule, sch, actor, phase, pen, trig, ev, sec) in OBLIGATIONS:
    exempt171 = "Y" if (not oid.startswith("LNK") and 4 <= sec < 17 and sec not in (8.1, 8.5) and oid != "OBL-SCP-03") else "N"
    obl_rows.append([oid, f"{dom} {dom_name[dom]}", title, req, act, rule, sch, actor,
                     PHASES[phase]["date"] if phase else "in force (other law)", pen, pen_txt.get(pen, ""),
                     "Y" if trig.get("always") else "N", trig_basis(trig), ",".join(trig.get("flags", [])), ",".join(trig.get("role", [])),
                     exempt171, "; ".join(ev), ", ".join(ctl_for.get(oid, []))])
wsO = sheet("Obligations", "Atomic Obligation Register (L1)",
      ["Obligation ID", "Domain", "Title", "Requirement", "Act / law ref", "Rule / instrument", "Schedule", "Actor", "In force from", "Penalty item", "Penalty cap", "Trigger: always", "Trigger: basis (any)", "Trigger: flags (all)", "Trigger: role", "Disapplied by s.17(1)?", "Evidence expected", "Controls"],
      obl_rows, [14, 26, 34, 70, 16, 18, 12, 14, 14, 9, 26, 9, 16, 22, 14, 11, 50, 30],
      note="Trigger columns drive the Applicability tab. s.17(1) keeps only s.8(1) accountability and s.8(5) security.")
OBL_N = len(obl_rows)

ctl_rows = [[c["id"], f"{c['domain']} {dom_name[c['domain']]}", c["title"], c["description"], c["type"], c["nature"], c["frequency"], c["owner_role"],
             ", ".join(c["obligations"]), c["test"], "; ".join(c["evidence"]), ", ".join(c["iso27001"]), ", ".join(c["iso27701"]), ", ".join(c["nist_csf"])] for c in CONTROLS]
sheet("Controls", "Control Library (L5) with crosswalk", ["Control ID", "Domain", "Title", "Description", "Type", "Nature", "Frequency", "Owner role", "Obligations", "Test procedure", "Evidence", "ISO 27001:2022", "ISO 27701:2019", "NIST CSF 2.0"],
      ctl_rows, [13, 26, 32, 60, 11, 14, 16, 18, 30, 50, 40, 16, 16, 18])

map_rows = []
for c in CONTROLS:
    for o in c["obligations"]:
        ob = [r for r in obl_rows if r[0] == o][0]
        map_rows.append([o, ob[2], c["id"], c["title"], ob[1], ob[9]])
sheet("Obl-Ctl Map", "Obligation to Control mapping (long form)", ["Obligation ID", "Obligation", "Control ID", "Control", "Domain", "Penalty item"], map_rows, [14, 40, 13, 40, 30, 10])

# ---------------------------------------------------------------- Legal
act_rows = [[f"s.{n}", t, ch, CHAPTERS[ch], ph, PHASES[ph]["date"], SECTION_PHASE_NOTES.get(n, ""), s, ", ".join(r)] for (n, t, ch, ph, s, r) in SECTIONS]
sheet("Act", "DPDP Act 2023 - 44 sections", ["Section", "Title", "Chapter", "Chapter title", "Phase", "Commences", "Split note", "Summary", "Rules"], act_rows, [8, 36, 8, 30, 7, 12, 24, 90, 18])
rs_rows = [[r, t, ph, PHASES[ph]["date"], ", ".join(f"s.{x}" for x in secs), s] for (r, t, ph, secs, s) in RULES] + [[c, t, ph, PHASES[ph]["date"], "", s] for (c, t, ph, s) in SCHEDULES]
sheet("Rules & Schedules", "DPDP Rules 2025 (23 rules) & Schedules", ["ID", "Title", "Phase", "Commences", "Sections", "Summary"], rs_rows, [9, 50, 7, 12, 14, 110])

# ---------------------------------------------------------------- Sectors
sec_rows, ret_rows, proc_rows = [], [], []
for s in SECTORS:
    sec_rows.append([f"SEC-{s['code']}", s["name"], s["covers"], ", ".join(s["regulators"]), "\n".join(f"- {a}: {b}" for a, b in s["laws"]), s["localisation"], "\n".join("- " + h for h in s["hotspots"]), "\n".join("- " + h for h in s["stuck"]), ", ".join(s["principals"])])
    for (rec, per, src, conf) in s["retention"]:
        ret_rows.append([s["name"], rec, per, src, conf])
    for p in s["processes"]:
        proc_rows.append([p[0], s["name"], p[1], p[2], "\n".join(p[3]), ", ".join(p[4]), ", ".join(p[5]), ", ".join(p[6]), ", ".join(p[7]), ", ".join(p[8]), ", ".join(p[9]), p[10]])
for p in COMMON:
    proc_rows.insert(0, [p[0], "All sectors (common)", p[1], p[2], "\n".join(p[3]), ", ".join(p[4]), ", ".join(p[5]), ", ".join(p[6]), ", ".join(p[7]), ", ".join(p[8]), ", ".join(p[9]), p[10]])
proc_rows.sort(key=lambda r: (r[1] != "All sectors (common)", r[0]))
ret_rows = [["All sectors", "Personal data, traffic data & processing logs", "Min. 1 year then erase", "DPDP Rule 8(3), Seventh Schedule", "high"],
            ["All sectors", "Logs & personal data for breach investigation", "1 year", "DPDP Rule 6(1)(e)", "high"],
            ["All sectors", "ICT system logs", "Rolling 180 days within India", "CERT-In Directions 2022 (iv)", "high"],
            ["All sectors", "Books of account (companies)", "8 years", "Companies Act 2013 s.128(5)", "high"],
            ["All sectors", "GST records", "72 months from due date of annual return", "CGST Act s.36", "high"],
            ["All sectors", "Consent Manager records", ">= 7 years", "DPDP First Schedule Part B", "high"]] + ret_rows
sheet("Sectors", "Sector Overlays (20 sectors)", ["ID", "Sector", "Covers", "Regulators", "Sector laws alongside DPDP", "Localisation / cross-border", "DPDP hotspots", "Where assessments get stuck", "Typical data principals"], sec_rows, [9, 24, 34, 26, 70, 40, 55, 50, 34])
sheet("Retention Anchors", "Retention anchors - reconcile s.8(7) erasure with 'retention required by law'", ["Sector", "Record", "Period", "Source", "Confidence"], ret_rows, [26, 44, 50, 40, 11], note="'verify' = confirm the current legal text before relying on it in a client deliverable.")
sheet("Process Catalogue", f"Process Catalogue - {len(proc_rows)} templates", ["Template ID", "Sector", "Department", "Process", "Typical activities", "Data principals", "Data categories", "Typical systems", "Typical third parties", "Typical basis", "Flags / context", "Assessor note"], proc_rows, [12, 22, 20, 32, 44, 26, 30, 26, 28, 16, 30, 40])
sheet("Data Elements", "Data Element Catalogue", ["ID", "Element", "Category", "Context tags", "Note"], [list(e) for e in ELEMENTS], [12, 40, 22, 20, 44])

# ---------------------------------------------------------------- Entity profile
wsE = wb.create_sheet("Entity Profile")
wsE["A1"] = "Entity Profile (L2) - one per legal entity"
wsE["A1"].font = TITLE
prof = [("Organisation ID", "ORG-DEMO"), ("Legal name", "DemoPay Finance Pvt Ltd (fictional)"), ("Sector(s)", "SEC-FIN"),
        ("Entity role", "data_fiduciary"), ("SDF notified? (Y/N)", "N"), ("Consent Manager? (Y/N)", "N"), ("Third Schedule class? (Y/N)", "N"),
        ("Website/app with DP interaction? (Y/N)", "Y"), ("Children exposure? (Y/N)", "N"), ("PwD guardian exposure? (Y/N)", "N"),
        ("Exemptions claimed", "none"), ("Regulators", "RBI, FIU-IND, CERT-In"), ("DPO / contact person", "Head - Privacy & Compliance"), ("Assessment cycle", "CYC-DEMO-2026Q4")]
for i, (k, v) in enumerate(prof, 3):
    wsE.cell(row=i, column=1, value=k).font = Font(name=F, bold=True, size=10)
    c = wsE.cell(row=i, column=2, value=v)
    c.font, c.fill, c.border = EXAMPLE, INPUT, BORDER
wsE.column_dimensions["A"].width = 40
wsE.column_dimensions["B"].width = 50
for r in (7, 8, 9, 10, 11, 12):
    dv_list(wsE, f"B{r}", LREF["YN"])
dv_list(wsE, "B6", LREF["Role"])
# Entity-level cells: B7 SDF, B8 CM, B9 Third schedule, B10 online

# ---------------------------------------------------------------- Activity register
AR_H = ["Activity ID", "Activity", "Department", "Process template", "Owner", "Purpose(s)", "Data principals", "Data elements", "Digital state", "Channel(s)", "Systems", "Third parties", "Retention rule", "Volume (approx. DPs)", "Confidence"] + [f"Basis: {b}" for b in BASIS_COLS] + [f"Flag: {f}" for f in FLAG_COLS] + ["Context tags", "Notes"]
ex = ["ACT-DEMO-CUS-001", "App onboarding & e-KYC", "Customer Onboarding", "FIN-01", "Head - Digital", "Onboarding & KYC", "Customer", "Name, PAN, Aadhaar, face, mobile", "digital", "Mobile app", "App backend; LOS; Salesforce", "AWS; KYC vendor; Salesforce; SMS aggregator", "RET-DEMO-001", 1200000, "high"]
bflags = {"consent": "Y", "s7d": "Y"}
fflags = {"processor": "Y", "cross_border": "Y", "online_presence": "Y", "legacy_data": "Y"}
ex += [bflags.get(b, "N") for b in BASIS_COLS] + [fflags.get(f, "N") for f in FLAG_COLS] + ["gov_id, biometric, financial", "Example row"]
N_ACT = 200
wsA = sheet("Activity Register", "Processing Activity Register (RoPA) - graph flattened to rows",
      AR_H, [ex], [16, 30, 18, 12, 16, 24, 18, 30, 11, 16, 26, 30, 12, 12, 11] + [8] * (len(BASIS_COLS) + len(FLAG_COLS)) + [22, 30],
      input_cols=range(len(AR_H)), blank_rows=N_ACT - 1, example_rows=1,
      note="One row per processing activity. Basis/flag columns take Y/N and feed the Applicability tab. Detailed graph (flows, events) lives in the vault.")
AR_START = 5
AR_END = AR_START + N_ACT - 1
b0 = 16  # column index (1-based) of first basis col
f0 = b0 + len(BASIS_COLS)
for j in range(b0, f0 + len(FLAG_COLS)):
    dv_list(wsA, f"{CL(j)}{AR_START}:{CL(j)}{AR_END}", LREF["YN"])
dv_list(wsA, f"I{AR_START}:I{AR_END}", LREF["DigitalState"])
dv_list(wsA, f"O{AR_START}:O{AR_END}", LREF["Confidence"])
BASIS_HDR = f"'Activity Register'!${CL(b0)}$4:${CL(f0-1)}$4"
FLAG_HDR = f"'Activity Register'!${CL(f0)}$4:${CL(f0+len(FLAG_COLS)-1)}$4"
# headers include prefix "Basis: " - use helper header row 3 with raw codes
for j, b in enumerate(BASIS_COLS):
    wsA.cell(row=3, column=b0 + j, value=b).font = Font(name=F, size=8, color="808080")
for j, f in enumerate(FLAG_COLS):
    wsA.cell(row=3, column=f0 + j, value=f).font = Font(name=F, size=8, color="808080")
BASIS_HDR = f"'Activity Register'!${CL(b0)}$3:${CL(f0-1)}$3"
FLAG_HDR = f"'Activity Register'!${CL(f0)}$3:${CL(f0+len(FLAG_COLS)-1)}$3"

# ---------------------------------------------------------------- Applicability matrix (formula engine)
wsX = wb.create_sheet("Applicability")
wsX["A1"] = "Applicability Matrix (L4) - Y = applies, N = not triggered, N/A-17(1) = disapplied by exemption, '-' = role not held"
wsX["A1"].font = TITLE
wsX["A2"] = "Computed from Obligations trigger columns x Activity Register Y/N columns x Entity Profile. The vault engine (Python) gives the same result with reasons."
wsX["A2"].font = Font(name=F, italic=True, size=9, color="595959")
NCOLS = 30
heads = ["Obligation ID", "Title", "Penalty", "Applies to # activities"]
for j, h in enumerate(heads, 1):
    c = wsX.cell(row=4, column=j, value=h)
    c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(wrap_text=True)
for k in range(NCOLS):
    col = 5 + k
    c = wsX.cell(row=4, column=col, value=f"=IF('Activity Register'!A{AR_START + k}=\"\",\"\",'Activity Register'!A{AR_START + k})")
    c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(text_rotation=90)
    wsX.column_dimensions[CL(col)].width = 5
wsX.row_dimensions[4].height = 110
wsX.column_dimensions["A"].width = 14
wsX.column_dimensions["B"].width = 38
wsX.column_dimensions["C"].width = 8
wsX.column_dimensions["D"].width = 11
for i in range(OBL_N):
    r = 5 + i
    orow = 5 + i  # Obligations sheet data row
    wsX.cell(row=r, column=1, value=f"=Obligations!A{orow}").font = BODY
    wsX.cell(row=r, column=2, value=f"=Obligations!C{orow}").font = BODY
    wsX.cell(row=r, column=3, value=f"=Obligations!J{orow}").font = BODY
    wsX.cell(row=r, column=4, value=f"=COUNTIF(E{r}:{CL(4 + NCOLS)}{r},\"Y\")").font = BODY
    for k in range(NCOLS):
        ar = AR_START + k
        brow = f"'Activity Register'!${CL(b0)}${ar}:${CL(f0-1)}${ar}"
        frow = f"'Activity Register'!${CL(f0)}${ar}:${CL(f0+len(FLAG_COLS)-1)}${ar}"
        req_b = f"Obligations!$M${orow}"
        req_f = f"Obligations!$N${orow}"
        role = f"Obligations!$O${orow}"
        actor = f"Obligations!$H${orow}"
        basis_ok = f"OR({req_b}=\"\",SUMPRODUCT(ISNUMBER(SEARCH(\",\"&{BASIS_HDR}&\",\",\",\"&{req_b}&\",\"))*({brow}=\"Y\"))>0)"
        nflags = f"IF({req_f}=\"\",0,LEN({req_f})-LEN(SUBSTITUTE({req_f},\",\",\"\"))+1)"
        flags_ok = f"SUMPRODUCT(ISNUMBER(SEARCH(\",\"&{FLAG_HDR}&\",\",\",\"&{req_f}&\",\"))*({frow}=\"Y\"))>={nflags}"
        role_ok = f"OR({role}=\"\",AND({role}=\"sdf\",'Entity Profile'!$B$7=\"Y\"),AND({role}=\"consent_manager\",'Entity Profile'!$B$8=\"Y\"))"
        actor_ok = f"AND(OR({actor}<>\"sdf\",'Entity Profile'!$B$7=\"Y\"),OR({actor}<>\"data_fiduciary\",'Entity Profile'!$B$6<>\"data_processor\"))"
        only_ex = f"AND('Activity Register'!${CL(b0+10)}${ar}=\"Y\",COUNTIF('Activity Register'!${CL(b0)}${ar}:${CL(b0+9)}${ar},\"Y\")=0)"
        only_ex2 = f"AND('Activity Register'!${CL(b0+11)}${ar}=\"Y\",COUNTIF('Activity Register'!${CL(b0)}${ar}:${CL(b0+10)}${ar},\"Y\")=0)"
        is_dpdp = f"LEFT($A{r},3)=\"OBL\""
        # entity flags OR activity flags for third_schedule / online / CM are handled by user setting activity flags; Entity Profile adds them:
        f_ts = FLAG_COLS.index("third_schedule"); f_op = FLAG_COLS.index("online_presence")
        ent_flag_bonus = f"(ISNUMBER(SEARCH(\"third_schedule\",{req_f}))*('Entity Profile'!$B$9=\"Y\")*('Activity Register'!${CL(f0+f_ts)}${ar}<>\"Y\")+ISNUMBER(SEARCH(\"online_presence\",{req_f}))*('Entity Profile'!$B$10=\"Y\")*('Activity Register'!${CL(f0+f_op)}${ar}<>\"Y\"))"
        flags_ok = f"SUMPRODUCT(ISNUMBER(SEARCH(\",\"&{FLAG_HDR}&\",\",\",\"&{req_f}&\",\"))*({frow}=\"Y\"))+{ent_flag_bonus}>={nflags}"
        formula = (f"=IF('Activity Register'!$A${ar}=\"\",\"\","
                   f"IF(AND({is_dpdp},{only_ex2},$A{r}<>\"OBL-SCP-03\",$A{r}<>\"OBL-RES-01\"),\"N/A-17(2)\","
                   f"IF(AND({is_dpdp},{only_ex},Obligations!$P${orow}=\"Y\"),\"N/A-17(1)\","
                   f"IF(NOT(AND({actor_ok},{role_ok})),\"-\","
                   f"IF(Obligations!$L${orow}=\"Y\",\"Y\","
                   f"IF(AND({basis_ok},{flags_ok}),\"Y\",\"N\"))))))")
        c = wsX.cell(row=r, column=5 + k, value=formula)
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(horizontal="center")
last = 4 + OBL_N
rng = f"E5:{CL(4 + NCOLS)}{last}"
wsX.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", fgColor="C6EFCE")))
wsX.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N"'], fill=PatternFill("solid", fgColor="F2F2F2")))
wsX.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(E5,3)="N/A"'], fill=PatternFill("solid", fgColor="FCE4D6")))
wsX.freeze_panes = "E5"

# ---------------------------------------------------------------- Control tests
CT_H = ["Test ID", "Cycle", "Control ID", "Control title", "Activity IDs", "Systems", "Test type", "Population", "Sample", "Tester", "Test date", "Evidence IDs", "Rating (0-4/NA)", "Result summary", "Finding IDs"]
ct_ex = ["TST-DEMO-CTL-NOT-01-01", "CYC-DEMO-2026Q4", "CTL-NOT-01", None, "ACT-DEMO-CUS-001", "App backend", "Inspection", "All app sign-up journeys", "3 journeys", "Assessor A", "2026-10-20", "EVD-DEMO-0003", "1", "Privacy-policy link only; no standalone itemised Rule 3 notice", "FND-DEMO-001"]
wsT = sheet("Control Tests", "Assessment Workbook - control tests (L6)", CT_H, [ct_ex], [22, 16, 12, 34, 22, 18, 13, 22, 10, 12, 11, 14, 10, 50, 14],
            input_cols=[0, 1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14], blank_rows=299, example_rows=1)
for r in range(5, 305):
    wsT.cell(row=r, column=4, value=f"=IFERROR(INDEX(Controls!$C$5:$C${4+len(CONTROLS)},MATCH(C{r},Controls!$A$5:$A${4+len(CONTROLS)},0)),\"\")").font = BODY
wsT.cell(row=4, column=16, value="Rating (numeric)").font = HDR_FONT
wsT.cell(row=4, column=16).fill = HDR_FILL
for r in range(5, 305):
    wsT.cell(row=r, column=16, value=f"=IF(OR(M{r}=\"\",M{r}=\"NA\"),\"\",VALUE(M{r}))").font = BODY
dv_list(wsT, "G5:G304", LREF["TestType"])
dv_list(wsT, "M5:M304", LREF["Rating"])

# ---------------------------------------------------------------- Findings
FH = ["Finding ID", "Title", "Cycle", "Department", "Activity ID", "System / flow", "Obligation IDs", "Control IDs", "Expected state", "Observed state", "Evidence IDs", "Penalty item", "Context (health/children/...)", "Likelihood (1-5)", "Impact (1-5)", "Risk score", "Severity", "Owner", "Remediation ID", "Target date", "Status", "Days to target"]
fex = ["FND-DEMO-004", "Shadow Excel with borrower & reference data", "CYC-DEMO-2026Q4", "Collections", "ACT-DEMO-CUS-003", "SYS-DEMO-XLS", "OBL-SEC-02, OBL-SEC-03", "CTL-SEC-02, CTL-SEC-04", "PD encrypted; access least-privilege", "Unencrypted file on a shared drive; 38 editors", "EVD-DEMO-0011", "P1", "financial", 4, 5, None, None, "Head - Collections", "REM-DEMO-004", "2027-02-28", "Accepted - remediation planned", None]
wsF = sheet("Findings", "Findings & Risk Register (L7)", FH, [fex], [14, 34, 16, 16, 18, 16, 22, 22, 32, 36, 14, 9, 16, 10, 10, 9, 10, 18, 14, 12, 22, 10],
            input_cols=[i for i in range(22) if i not in (15, 16, 21)], blank_rows=299, example_rows=1)
for r in range(5, 305):
    wsF.cell(row=r, column=16, value=f"=IF(OR(N{r}=\"\",O{r}=\"\"),\"\",N{r}*O{r})").font = BODY
    wsF.cell(row=r, column=17, value=f"=IF(P{r}=\"\",\"\",IF(P{r}>=17,\"Critical\",IF(P{r}>=10,\"High\",IF(P{r}>=5,\"Medium\",\"Low\"))))").font = BODY
    wsF.cell(row=r, column=22, value=f"=IF(OR(T{r}=\"\",U{r}=\"Closed\"),\"\",IF(ISNUMBER(T{r}),T{r},DATEVALUE(T{r}))-TODAY())").font = BODY
dv_list(wsF, "N5:O304", LREF["Scale"])
dv_list(wsF, "L5:L304", LREF["Penalty"])
dv_list(wsF, "U5:U304", LREF["FindingStatus"])
for band, col in (("Critical", "F8CBAD"), ("High", "FCE4D6"), ("Medium", "FFF2CC"), ("Low", "E2EFDA")):
    wsF.conditional_formatting.add("Q5:Q304", CellIsRule(operator="equal", formula=[f'"{band}"'], fill=PatternFill("solid", fgColor=col)))
wsF["P4"].comment = Comment("Likelihood x Impact. Bands: 17-25 Critical, 10-16 High, 5-9 Medium, 1-4 Low (see vault: Risk Methodology).", "DUATF")

sheet("Remediation", "Remediation Plan", ["Remediation ID", "Finding IDs", "Workstream", "Action", "Owner", "Start", "Target date", "Status", "Retest test ID", "Notes"],
      [["REM-DEMO-004", "FND-DEMO-004", "Security", "Migrate tracker into CRM collections module; delete Excel; DLP rule for PD on shared drives", "Head - Collections", "2026-11-01", "2027-02-28", "Not started", "", ""]],
      [14, 16, 16, 60, 18, 12, 12, 14, 16, 30], input_cols=range(10), blank_rows=299, example_rows=1)

# ---------------------------------------------------------------- Rights log
RH = ["Request ID", "Received on", "Channel", "Right", "Identifier used", "Identity verified (Y/N)", "Systems searched", "Processors notified", "Legal retention applied", "Internal SLA (days)", "Due by", "Statutory max (90d)", "Closed on", "Days taken", "Outcome", "Breached internal SLA?"]
rex = ["DSR-2027-0001", "2027-05-20", "Website", "erasure", "mobile", "Y", "App, LOS, CRM", "Salesforce, SMS aggregator", "KYC & loan records kept 5 yrs (PMLA)", 30, None, None, "2027-06-10", None, "Partially erased - explained", None]
wsR = sheet("Rights Log", "Data Principal Rights & Grievance Log (R14 - max 90 days)", RH, [rex], [16, 12, 14, 16, 14, 10, 26, 24, 30, 10, 12, 12, 12, 9, 30, 10],
            input_cols=[0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 12, 14], blank_rows=499, example_rows=1)
for r in range(5, 505):
    wsR.cell(row=r, column=11, value=f"=IF(B{r}=\"\",\"\",IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r}))+J{r})").number_format = "yyyy-mm-dd"
    wsR.cell(row=r, column=12, value=f"=IF(B{r}=\"\",\"\",IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r}))+90)").number_format = "yyyy-mm-dd"
    wsR.cell(row=r, column=14, value=f"=IF(OR(B{r}=\"\",M{r}=\"\"),\"\",IF(ISNUMBER(M{r}),M{r},DATEVALUE(M{r}))-IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r})))")
    wsR.cell(row=r, column=16, value=f"=IF(B{r}=\"\",\"\",IF(M{r}=\"\",IF(TODAY()>IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r}))+J{r},\"OVERDUE\",\"open\"),IF(N{r}>J{r},\"Y\",\"N\")))")
    for col in (11, 12, 14, 16):
        wsR.cell(row=r, column=col).font = BODY
dv_list(wsR, "D5:D504", LREF["Right"])
dv_list(wsR, "F5:F504", LREF["YN"])
wsR.conditional_formatting.add("P5:P504", CellIsRule(operator="equal", formula=['"OVERDUE"'], fill=PatternFill("solid", fgColor="F8CBAD")))
wsR["B4"].comment = Comment("Enter dates as text yyyy-mm-dd (e.g. 2027-05-20).", "DUATF")

# ---------------------------------------------------------------- Breach register
BH = ["Breach ID", "Aware at (yyyy-mm-dd hh:mm)", "Description", "PD affected (Y/N)", "# DPs affected", "Data elements", "Systems", "Processor involved", "CERT-In due (6h)", "CERT-In reported at", "Board initial sent at", "Board 72h due", "Board 72h report at", "Sector regulator", "Sector reported at", "DPs notified at", "Root cause", "Status"]
bex = ["BRE-2027-001", "2027-06-02 10:30", "Misconfigured storage bucket exposed KYC images", "Y", 1800, "Photo, PAN, Aadhaar (masked)", "App backend", "KYC vendor", None, "2027-06-02 14:10", "2027-06-02 16:00", None, "2027-06-04 18:00", "RBI", "2027-06-02 15:00", "2027-06-03 12:00", "Public ACL set during vendor migration", "closed"]
wsB = sheet("Breach Register", "Breach Register - one awareness timestamp drives every clock", BH, [bex], [14, 18, 34, 9, 10, 24, 16, 16, 18, 18, 18, 18, 18, 14, 18, 18, 30, 10],
            input_cols=[i for i in range(18) if i not in (8, 11)], blank_rows=99, example_rows=1)
for r in range(5, 105):
    wsB.cell(row=r, column=9, value=f"=IF(B{r}=\"\",\"\",IF(ISNUMBER(B{r}),B{r},DATEVALUE(LEFT(B{r},10))+TIMEVALUE(MID(B{r},12,5)))+6/24)").number_format = "yyyy-mm-dd hh:mm"
    wsB.cell(row=r, column=12, value=f"=IF(B{r}=\"\",\"\",IF(ISNUMBER(B{r}),B{r},DATEVALUE(LEFT(B{r},10))+TIMEVALUE(MID(B{r},12,5)))+3)").number_format = "yyyy-mm-dd hh:mm"
    wsB.cell(row=r, column=9).font = BODY
    wsB.cell(row=r, column=12).font = BODY
wsB["I4"].comment = Comment("CERT-In Directions 28.04.2022: report within 6 hours of noticing. Sector regulators may be stricter (e.g. RBI).", "DUATF")
wsB["L4"].comment = Comment("DPDP Rule 7(2)(b): detailed report within 72 hours of becoming aware (or longer if the Board allows on a written request). The initial intimation is 'without delay'.", "DUATF")

# ---------------------------------------------------------------- Third parties
TH = ["TP ID", "Name", "Type", "Role (processor / DF / joint)", "Role rationale", "Services", "Data received", "Activities", "Country / support location", "Sub-processors", "DPA status", "R6 security clause (Y/N)", "Breach SLA (hours)", "Erasure on exit (Y/N)", "Audit rights (Y/N)", "Certifications", "Risk tier", "Last review"]
tex = ["TP-DEMO-COLL", "Field collection agency", "Collection agency", "data_processor", "Acts only on DemoPay allocation", "Field recovery", "Name, mobile, address, dues", "ACT-DEMO-CUS-003", "India", "Unknown", "none", "N", "", "N", "N", "", "High", ""]
wsV = sheet("Third Parties", "Third-Party / Processor Register (D13)", TH, [tex], [14, 26, 20, 16, 28, 20, 28, 20, 16, 16, 16, 9, 9, 9, 9, 16, 9, 12], input_cols=range(18), blank_rows=299, example_rows=1)
dv_list(wsV, "C5:C304", LREF["TPType"])
dv_list(wsV, "K5:K304", LREF["DPAStatus"])
sheet("Transfers", "Cross-Border Transfer Register (D14)", ["Transfer ID", "Activity ID", "From system", "Recipient (TP ID)", "Country", "Transfer type", "Data elements", "Purpose", "Sector localisation check", "On restricted list? (as notified)", "Safeguards", "Last checked"],
      [["XB-DEMO-001", "ACT-DEMO-MKT-001", "SYS-DEMO-CRM", "TP-DEMO-SF", "USA", "cross-border to processor", "Name, mobile, email, segment", "Marketing", "RBI payment data NOT included - OK", "No list notified yet", "Vendor DPA, encryption, SSO", "2026-10-20"]],
      [12, 18, 16, 16, 12, 22, 28, 16, 30, 16, 30, 12], input_cols=range(12), blank_rows=199, example_rows=1)

# ---------------------------------------------------------------- Evidence (PBC)
pbc = []
for line in PB.PBC_LIST.splitlines():
    m = re.match(r"\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*(D\d+(?:/D\d+)?)\s*\|\s*(.+?)\s*\|", line)
    if m:
        pbc.append([f"PBC-{int(m.group(1)):02d}", m.group(2), m.group(3), m.group(4), "", "requested", "", "", "", ""])
sheet("Evidence (PBC)", "Evidence Request List & Tracker", ["PBC ID", "Evidence", "Domain", "Typical owner", "Client owner", "Status", "Received on", "Valid until", "File link", "Notes"], pbc,
      [9, 60, 10, 20, 18, 12, 12, 12, 30, 30], input_cols=[4, 5, 6, 7, 8, 9], blank_rows=50)
dv_list(wb["Evidence (PBC)"], "F5:F100", LREF["EvidenceStatus"])

# ---------------------------------------------------------------- Dashboard (first tab)
wsD = wb.create_sheet("Dashboard", 1)
wsD["A1"] = "DPDP Readiness Dashboard"
wsD["A1"].font = TITLE
wsD["A2"] = "All figures are formulas over the working tabs."
wsD["A2"].font = Font(name=F, italic=True, size=9, color="595959")
k = [
 ("Days to Phase 2 (13 Nov 2026)", "=DATE(2026,11,13)-TODAY()"),
 ("Days to Phase 3 (13 May 2027)", "=DATE(2027,5,13)-TODAY()"),
 ("Processing activities recorded", f"=COUNTA('Activity Register'!A{AR_START}:A{AR_END})"),
 ("  of which low-confidence mapping", f"=COUNTIF('Activity Register'!O{AR_START}:O{AR_END},\"low\")"),
 ("  of which involve children", f"=COUNTIF('Activity Register'!{CL(f0+FLAG_COLS.index('children'))}{AR_START}:{CL(f0+FLAG_COLS.index('children'))}{AR_END},\"Y\")"),
 ("  of which cross-border", f"=COUNTIF('Activity Register'!{CL(f0+FLAG_COLS.index('cross_border'))}{AR_START}:{CL(f0+FLAG_COLS.index('cross_border'))}{AR_END},\"Y\")"),
 ("Obligation x activity pairs applicable (first 30 activities)", f"=COUNTIF(Applicability!E5:{CL(4+NCOLS)}{last},\"Y\")"),
 ("Controls tested", "=COUNTA('Control Tests'!C5:C304)"),
    ("Average control rating (0-4)", "=IFERROR(AVERAGE('Control Tests'!P5:P304),0)"),
 ("Readiness index (avg rating / 4)", "=B12/4"),
 ("Open findings - Critical", "=COUNTIFS(Findings!Q5:Q304,\"Critical\",Findings!U5:U304,\"<>Closed\")"),
 ("Open findings - High", "=COUNTIFS(Findings!Q5:Q304,\"High\",Findings!U5:U304,\"<>Closed\")"),
 ("Open findings - Medium", "=COUNTIFS(Findings!Q5:Q304,\"Medium\",Findings!U5:U304,\"<>Closed\")"),
 ("Open findings - Low", "=COUNTIFS(Findings!Q5:Q304,\"Low\",Findings!U5:U304,\"<>Closed\")"),
 ("Findings past target date", "=COUNTIFS(Findings!V5:V304,\"<0\")"),
 ("Rights requests logged", "=COUNTA('Rights Log'!A5:A504)"),
 ("Rights requests overdue (internal SLA)", "=COUNTIF('Rights Log'!P5:P504,\"OVERDUE\")"),
 ("Breaches logged", "=COUNTA('Breach Register'!A5:A104)"),
 ("Third parties recorded", "=COUNTA('Third Parties'!A5:A304)"),
 ("  processors with no DPA", "=COUNTIFS('Third Parties'!D5:D304,\"*processor*\",'Third Parties'!K5:K304,\"none\")"),
 ("Evidence items outstanding", "=COUNTIF('Evidence (PBC)'!F5:F100,\"requested\")+COUNTIF('Evidence (PBC)'!F5:F100,\"insufficient\")"),
]
for i, (lab, f) in enumerate(k, 4):
    wsD.cell(row=i, column=1, value=lab).font = Font(name=F, size=10, bold=not lab.startswith("  "))
    c = wsD.cell(row=i, column=2, value=f)
    c.font = Font(name=F, size=11, bold=True, color="1F3864")
    c.border = BORDER
wsD["B12"].number_format = "0.00"
wsD["B13"].number_format = "0%"
wsD.column_dimensions["A"].width = 55
wsD.column_dimensions["B"].width = 16
# obligations per domain summary
wsD["D3"] = "Library size by domain"
wsD["D3"].font = Font(name=F, bold=True, size=11)
for j, h in enumerate(["Domain", "Obligations", "Controls", "Rs 200-250 cr tier obligations"], 4):
    c = wsD.cell(row=4, column=j, value=h)
    c.font, c.fill = HDR_FONT, HDR_FILL
for i, (did, name, _) in enumerate(DOMAINS, 5):
    wsD.cell(row=i, column=4, value=f"{did} {name}").font = BODY
    wsD.cell(row=i, column=5, value=f"=COUNTIF(Obligations!$B$5:$B${4+OBL_N},D{i})").font = BODY
    wsD.cell(row=i, column=6, value=f"=COUNTIF(Controls!$B$5:$B${4+len(CONTROLS)},D{i})").font = BODY
    wsD.cell(row=i, column=7, value=f"=COUNTIFS(Obligations!$B$5:$B${4+OBL_N},D{i},Obligations!$J$5:$J${4+OBL_N},\"P1\")+COUNTIFS(Obligations!$B$5:$B${4+OBL_N},D{i},Obligations!$J$5:$J${4+OBL_N},\"P2\")+COUNTIFS(Obligations!$B$5:$B${4+OBL_N},D{i},Obligations!$J$5:$J${4+OBL_N},\"P3\")").font = BODY
wsD.column_dimensions["D"].width = 42
wsD.column_dimensions["E"].width = 12
wsD.column_dimensions["F"].width = 10
wsD.column_dimensions["G"].width = 16

order = ["README", "Dashboard", "Entity Profile", "Activity Register", "Applicability", "Control Tests", "Findings", "Remediation", "Rights Log", "Breach Register",
         "Third Parties", "Transfers", "Evidence (PBC)", "Obligations", "Controls", "Obl-Ctl Map", "Act", "Rules & Schedules", "Sectors", "Retention Anchors", "Process Catalogue", "Data Elements", "Lists"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = "FFC000" if ws.title in order[2:13] else ("1F3864" if ws.title in ("README", "Dashboard") else "8EA9DB")
wb.save(OUT)
print("saved", OUT)
